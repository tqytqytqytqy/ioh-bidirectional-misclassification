"""Verify the released aggregate snapshot; does not refit source-data analyses."""
from pathlib import Path
import csv
import hashlib
import json
import math
import sys
from collections import defaultdict
from zipfile import ZipFile
import openpyxl

ROOT = Path(__file__).resolve().parents[2]

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def verify():
    checks = {}
    plan = json.loads((ROOT/'current/presentation/workbook_plan.json').read_text())
    layout = json.loads((ROOT/'current/provenance/workbook_layout.json').read_text())
    book = openpyxl.load_workbook(ROOT/'current/workbook/10_SDC2_Consolidated_Workbook.xlsx', data_only=False)
    assert book.sheetnames == [s['name'] for s in plan] and len(book.sheetnames) == 25
    cells = numbers = 0
    for spec, page in zip(plan, layout):
        sheet = book[spec['name']]
        assert sheet.sheet_state == 'visible'
        for block, section in zip(spec['blocks'], page['sections']):
            for r, row in enumerate(block['rows']):
                vals = [v for v in row if v not in (None, '')]
                merged = len(vals) == 1 and isinstance(vals[0], str)
                merged |= r in section['repeatedNoteRows']
                expected = [vals[0]] + [None]*(len(row)-1) if merged else row
                for c, value in enumerate(expected):
                    actual = sheet.cell(section['headerRow']+r, c+1).value
                    if value == '': value = None
                    if isinstance(value, (int, float)):
                        assert isinstance(actual, (int, float)) and math.isclose(actual, value, rel_tol=1e-13, abs_tol=1e-13)
                        numbers += 1
                    else:
                        assert actual == value, (sheet.title, section['headerRow']+r, c+1)
                    cells += 1
        for row in sheet:
            for cell in row:
                assert cell.data_type not in ('f', 'e'), (sheet.title, cell.coordinate)
    checks['workbook'] = {'sheets':25, 'logical_cells':cells, 'numeric_cells':numbers}
    records = json.loads((ROOT/'current/provenance/source_file_map.json').read_text())
    for item in records:
        assert sha(ROOT/item['target']) == item['sha256'], item['target']
    checks['source_map_files'] = len(records)
    recovery = ROOT/'source_layout/EJA_审稿意见逐项修订版_20261016/aggregate/recovery_summary.csv'
    with recovery.open() as stream: rows = list(csv.DictReader(stream))
    groups = defaultdict(dict)
    for row in rows:
        key = tuple(row[k] for k in ('scenario','normal_sec','interval_min','horizon_sec','unit','weight'))
        groups[key][row['metric']] = float(row['estimate'])
    for values in groups.values():
        assert math.isclose(sum(values[k] for k in ('persistent_low','display_cleared','reference_recurred','unknown')), 1, abs_tol=1e-10)
        assert math.isclose(values['persistent_lower'],values['persistent_low'],abs_tol=1e-10)
        assert math.isclose(values['persistent_upper'],values['persistent_low']+values['unknown'],abs_tol=1e-10)
    primary = [r for r in rows if r['scenario']=='original_all' and r['normal_sec']=='20' and float(r['interval_min'])==5 and r['horizon_sec']=='120' and r['unit']=='distinct_recovery' and r['weight']=='unit_equal' and r['metric']=='persistent_low']
    assert len(primary)==1 and int(primary[0]['units'])==164 and int(primary[0]['patients'])==78
    assert round(float(primary[0]['estimate'])*100,1)==25.6
    checks['recovery_partitions'] = len(groups)
    summary = ROOT/'source_layout/跨库分辨率与异质性分析_20261015/aggregate/summary.csv'
    with summary.open() as stream: rows = list(csv.DictReader(stream))
    groups = defaultdict(dict)
    for row in rows:
        key = tuple(row[k] for k in ('source','resolution','scenario','threshold','interval_min','event_min_sec','estimand'))
        groups[key][row['metric']] = float(row['estimate']) if row['estimate'] else float('nan')
    identities = 0
    for values in groups.values():
        if all(k in values and math.isfinite(values[k]) for k in ('hidden_ratio','overdisplay_ratio','net_ratio')):
            assert math.isclose(values['net_ratio'], values['overdisplay_ratio']-values['hidden_ratio'],abs_tol=1e-10)
            identities += 1
    checks['cross_database_area_identities'] = identities
    blocked = {'subject_id','patient_id','caseid','case_id','record_id','hadm_id','stay_id','track_id'}
    csvs = 0
    for path in (ROOT/'source_layout').rglob('*.csv'):
        with path.open() as stream: fields = next(csv.reader(stream), [])
        assert not blocked.intersection(x.lower() for x in fields), str(path.relative_to(ROOT))
        csvs += 1
    for path in (ROOT/'current').rglob('*'):
        if path.suffix in ('.xlsx','.pptx'):
            with ZipFile(path) as package:
                for name in package.namelist():
                    assert 'externalLinks/' not in name and '/embeddings/' not in name
                    if name.endswith('.xml'):
                        body=package.read(name)
                        assert b'/Users/' not in body and b'/Volumes/' not in body, name
    checks['aggregate_csv_header_checks'] = csvs
    manifest = ROOT/'current/provenance/outputs_manifest.json'
    if manifest.exists():
        data = json.loads(manifest.read_text())
        for entry in data['files']: assert sha(ROOT/entry['path'])==entry['sha256'],entry['path']
        checks['manifest_files'] = len(data['files'])
    checks['scope'] = 'Released aggregate values, source provenance and package structure only; no raw-data rerun or clinical adjudication.'
    return {'status':'PASS', 'checks':checks}

if __name__ == '__main__':
    result=verify()
    if len(sys.argv)>1: Path(sys.argv[1]).write_text(json.dumps(result,indent=2)+'\n')
    print(json.dumps(result,indent=2))
